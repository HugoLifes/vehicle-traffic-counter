#!/usr/bin/env python3
"""
Compara lo que cuenta el sistema contra un aforo real medido en campo.

Esta es la unica herramienta que da una cifra de exactitud *absoluta*. Todo
lo demas que se ha medido en el proyecto es relativo ("cuenta un 43 % mas
que antes"), y un informe de aforo tiene que declarar una exactitud.

En referencias/aforo_real/ hay DOS mediciones de campo del mismo tramo y el
mismo dia (19-ago-2026), y no coinciden entre si:

- `conteo_manual_24h.xlsx` — aforo contado por una persona, 24 h, por
  cuartos de hora, por sentido y por clase. **Es la referencia.**
- `miguel_de_la_madrid_*.xls` — contador de ejes, uno por sentido.

Contrastados entre si, el tubo perdio el 31 % del transito en el sentido
pte-ote (1 916 contra 2 777 contados a mano) mientras acertaba en ote-pte
(0.95x). Por eso `--fuente auto` toma el conteo manual cuando existe: una
persona contando es la referencia, y el tubo es un instrumento que aqui
fallo en un sentido.

Esto importa porque durante un tiempo se dio por bueno el tubo y se
concluyo que nuestro sistema sobrecontaba la calzada del fondo en un 36 %.
Contra el conteo manual esa misma calzada sale en 0.94x: no sobrecontaba,
el tubo subcontaba.

Uso:
    python tools/comparar_aforo_real.py --proyecto 2
    python tools/comparar_aforo_real.py --proyecto 2 --fuente tubo
    python tools/comparar_aforo_real.py --proyecto 2 --salida data/comparacion.csv
"""
from __future__ import annotations

import argparse
import csv
import re
import sqlite3
import statistics
import sys
from collections import defaultdict
from pathlib import Path

RAIZ = Path(__file__).resolve().parent.parent
REFERENCIAS = RAIZ / "referencias" / "aforo_real"
BD = RAIZ / "data" / "traffic.db"


# Sentidos como los escribe la empresa. El primer formato solo traia
# OTE-PTE y PTE-OTE, y el lector los tenia fijos: un aforo de una calle
# norte-sur, o escrito "NORTE-SUR", no se leia y la herramienta decia que no
# habia conteo manual.
_RUMBO = r"(?:NTE|SUR|OTE|PTE|NORTE|ORIENTE|PONIENTE)"
PATRON_SENTIDO = re.compile(rf"{_RUMBO}\s*-\s*{_RUMBO}", re.I)
# La fecha viene en el titulo de la hoja: "BLVD MIGUEL DE LA MADRID
# 19-AGOSTO-2026". Es el candado contra comparar un video con el aforo de
# otro dia, que en la carpeta de referencias conviven.
PATRON_FECHA = re.compile(r"(\d{1,2})\s*-\s*([A-Za-zÁÉÍÓÚáéíóú]+)\s*-\s*(\d{4})")
MESES = {"ENERO": 1, "FEBRERO": 2, "MARZO": 3, "ABRIL": 4, "MAYO": 5,
         "JUNIO": 6, "JULIO": 7, "AGOSTO": 8, "SEPTIEMBRE": 9, "SETIEMBRE": 9,
         "OCTUBRE": 10, "NOVIEMBRE": 11, "DICIEMBRE": 12}
CLASES_MANUAL = ("A", "B", "C", "T-S", "T-S-R")


def _fecha_de_titulo(txt: str) -> str | None:
    m = PATRON_FECHA.search(txt or "")
    if not m:
        return None
    mes = MESES.get(m.group(2).upper())
    return f"{int(m.group(3)):04d}-{mes:02d}-{int(m.group(1)):02d}" if mes else None


def leer_conteo_manual_clases(ruta: Path):
    """Aforo contado por una persona, con su clase.

    Devuelve (fecha, {sentido: {minuto: {clase: vehiculos}}}).

    Un archivo trae los DOS sentidos, una hoja cada uno, y cada hoja se
    parte en dos bloques lado a lado: la mitad AM a la izquierda y la PM a
    la derecha, con cinco clases (A, B, C, T-S, T-S-R) por bloque. La
    columna de la hora se localiza por su encabezado 'Hr/Mov' en vez de
    fijarla, porque no cae en la misma letra en las dos hojas; y el nombre
    de cada clase se lee de la fila de abajo en vez de suponer el orden.

    El formato NO trae columna de motocicletas. O van dentro de A o no se
    cuentan, y eso hay que preguntarlo: nosotros si las contamos.
    """
    import openpyxl  # solo aqui: el resto de la herramienta no lo necesita

    libro = openpyxl.load_workbook(str(ruta), data_only=True)
    real: dict[str, dict[int, dict[str, float]]] = {}
    fecha = None

    for hoja in libro.worksheets:
        columnas_hora, sentido = [], None
        for fila in hoja.iter_rows(min_row=1, max_row=6):
            for celda in fila:
                txt = str(celda.value or "").strip()
                if txt.lower().startswith("hr/mov"):
                    columnas_hora.append((celda.row, celda.column))
                elif PATRON_SENTIDO.fullmatch(txt):
                    sentido = re.sub(r"\s+", "", txt).lower()
                elif fecha is None:
                    fecha = _fecha_de_titulo(txt)
        if not columnas_hora:
            continue
        # Sin un sentido reconocible se usa el nombre de la hoja: mejor un
        # nombre raro que descartar la hoja entera en silencio.
        sentido = sentido or hoja.title.strip().lower()

        por_minuto: dict[int, dict[str, float]] = {}
        for fila_hr, col in columnas_hora:
            nombres = []
            for i in range(1, 6):
                n = str(hoja.cell(fila_hr + 1, col + i).value or "").strip().upper()
                nombres.append(n if n in CLASES_MANUAL else CLASES_MANUAL[i - 1])
            for fila in range(fila_hr + 1, hoja.max_row + 1):
                minuto = _minuto_de_rango(hoja.cell(fila, col).value)
                if minuto is None or minuto in por_minuto:
                    continue
                valores = [hoja.cell(fila, col + i).value for i in range(1, 6)]
                # Una fila con TODAS las clases en blanco no es "cero
                # vehiculos": es un cuarto de hora que nadie conto. El conteo
                # de 24 h venia lleno y no importaba; uno de tres franjas deja
                # el resto en blanco, y leido como cero la comparacion ponia
                # trece horas nuestras contra ceros.
                if all(not isinstance(v, (int, float)) for v in valores):
                    continue
                por_minuto[minuto] = {
                    nombre: float(v) if isinstance(v, (int, float)) else 0.0
                    for nombre, v in zip(nombres, valores)}
        if por_minuto:
            real[sentido] = por_minuto
    return fecha, real


def leer_conteo_manual(ruta: Path) -> dict[str, dict[int, float]]:
    """Aforo contado por una persona: {sentido: {minuto: vehiculos}}."""
    _, real = leer_conteo_manual_clases(ruta)
    return {s: {m: sum(c.values()) for m, c in d.items()} for s, d in real.items()}


def _minuto_de_rango(etiqueta) -> int | None:
    """'7:15 AM - 7:30 AM' -> 435, el minuto del dia en que empieza."""
    m = re.match(r"\s*(\d{1,2}):(\d{2})\s*(AM|PM)", str(etiqueta or ""), re.I)
    if not m:
        return None
    h, mi, ap = int(m.group(1)), int(m.group(2)), m.group(3).upper()
    if ap == "PM" and h != 12:
        h += 12
    if ap == "AM" and h == 12:
        h = 0
    return h * 60 + mi


def cargar_referencia(carpeta: Path, fuente: str = "auto", fecha: str | None = None):
    """El conteo manual manda sobre el contador de ejes cuando ambos existen.

    Contrastados entre si sobre este mismo tramo y dia, el tubo perdio el
    31 % del transito en un sentido (1 916 contra 2 777 contados a mano)
    mientras acertaba en el otro (0.95x). Una persona contando es la
    referencia; el tubo es un instrumento que puede fallar y aqui fallo.

    Devuelve (totales, clases, fecha). `clases` va en la taxonomia de la
    empresa (el tubo se traduce de sus clases por ejes); `fecha` es None si
    no se pudo leer. `fecha` escoge el dia cuando el reporte trae varios.
    """
    manuales = sorted(carpeta.glob("*.xlsx"))
    tubos = sorted(carpeta.glob("*.xls"))

    if fuente in ("auto", "manual") and manuales:
        real: dict[str, dict[int, float]] = {}
        clases: dict[str, dict[int, dict[str, float]]] = {}
        fechas = set()
        origen: dict[str, str] = {}
        for ruta in manuales:
            fecha, por_clase = leer_conteo_manual_clases(ruta)
            if fecha:
                fechas.add(fecha)
            for sentido, datos in por_clase.items():
                # Antes el segundo archivo pisaba al primero sin decir nada.
                # Con el aforo viejo y el nuevo en la misma carpeta, eso
                # compara un sentido de un dia contra el otro sentido de otro.
                if sentido in origen:
                    sys.exit(f"El sentido {sentido} viene en dos archivos "
                             f"({origen[sentido]} y {ruta.name}). Deja en "
                             f"{carpeta} solo el conteo de ESTE video, o usa "
                             "--referencias con una carpeta para cada aforo.")
                origen[sentido] = ruta.name
                clases[sentido] = datos
                real[sentido] = {m: sum(c.values()) for m, c in datos.items()}
                print(f"  {ruta.name} [conteo manual]: sentido {sentido}, "
                      f"{fecha or 'fecha sin leer'}, {len(datos)}/96 cuartos de "
                      f"hora, {int(sum(real[sentido].values()))} vehiculos")
        if len(fechas) > 1:
            sys.exit(f"Los conteos de {carpeta} son de dias distintos: "
                     f"{', '.join(sorted(fechas))}.")
        if real:
            return real, clases, (fechas.pop() if fechas else None)
    if fuente == "manual":
        sys.exit(f"No hay conteo manual (.xlsx) en {carpeta}")

    if not tubos:
        sys.exit(f"No hay reportes de referencia en {carpeta}")
    return cargar_tubos(tubos, fecha)


def cargar_tubos(tubos, fecha):
    """Contador de ejes: (totales, clases SCT, fecha) del dia `fecha`.

    Lo lee `contador_ejes.py`, que separa las secciones del reporte (un
    archivo trae la clasificacion y las velocidades de DOS carriles) y los
    dias (el del frontal trae cinco). Se usa el carril 1: el 2 es un canal
    sin mangueras, casi en cero, que no correlaciona con el transito.

    Cuando del mismo sentido llegan el clasificatorio y el de volumen, manda
    el clasificatorio (trae las clases) y el de volumen se usa para
    comprobar que el aparato cuadra consigo mismo.
    """
    sys.path.insert(0, str(RAIZ / "tools"))
    import contador_ejes as ce

    porsentido: dict[str, dict] = {}
    for ruta in tubos:
        meta, secciones, _ = ce.leer_todo(str(ruta))
        sentido = meta.get("sentido", ruta.stem.lower())
        tipo = meta.get("tipo", "clasificatorio")
        datos = secciones.get((tipo, 1), {})
        dias = sorted(datos)
        dia = fecha if fecha in datos else (dias[0] if len(dias) == 1 else None)
        if dia is None:
            sys.exit(f"{ruta.name} trae {len(dias)} dias ({', '.join(dias)}) y ninguno "
                     f"es el del video ({fecha}).")
        porsentido.setdefault(sentido, {})[tipo] = (ruta.name, meta, datos[dia])

    real, clases = {}, {}
    for sentido, tipos in sorted(porsentido.items()):
        if "clasificatorio" in tipos:
            nombre, meta, d = tipos["clasificatorio"]
            clases[sentido] = {m: _clases_sct(v, ce.SCT) for m, v in d.items()}
            real[sentido] = {m: sum(v.values()) for m, v in d.items()}
        else:
            nombre, meta, d = tipos["volumen"]
            real[sentido] = dict(d)
        print(f"  {nombre} [contador de ejes, serie {meta.get('serie', '?')}, mangueras a "
              f"{meta.get('separacion', '?')}]: sentido {sentido}, {len(real[sentido])} cuartos "
              f"de hora, {int(sum(real[sentido].values()))} vehiculos")
        if "clasificatorio" in tipos and "volumen" in tipos:
            vol = tipos["volumen"][2]
            comunes = set(vol) & set(real[sentido])
            a = sum(real[sentido][m] for m in comunes)
            b = sum(vol[m] for m in comunes)
            print(f"    el reporte de volumen del mismo aparato da {int(b)} contra {int(a)}"
                  f" ({100 * (a - b) / b:+.1f} %)" if b else "")
    return real, (clases or None), fecha


def _clases_sct(v, sct):
    out: dict[str, float] = defaultdict(float)
    for clase, n in v.items():
        out[sct.get(clase, "OTRO")] += n
    return dict(out)


def diagnosticar_sin_cruces(bd: Path, proyecto: int) -> str:
    """Un proyecto sin cruces casi siempre es un proyecto que todavia no ha
    contado, no un error. Se dice cual es el caso y que sigue, en vez de
    dejar al usuario mirando un "no tiene cruces"."""
    con = _conectar_ro(bd)
    estados = dict(
        con.execute(
            "select status, count(*) from video_jobs where project_id=? group by status",
            (proyecto,),
        ).fetchall()
    )
    con.close()
    if not estados:
        return f"El proyecto {proyecto} no tiene videos."
    resumen = ", ".join(f"{n} {e}" for e, n in sorted(estados.items()))
    if estados.get("awaiting_calibration"):
        return (
            f"El proyecto {proyecto} todavia no ha contado nada ({resumen}).\n"
            "Revisa el encuadre en la plataforma y presiona 'Empezar conteo'.\n"
            "Vuelve a correr esto cuando los videos esten en 'listo'."
        )
    if estados.get("queued") or estados.get("processing"):
        return f"El proyecto {proyecto} sigue contando ({resumen}). Espera a que termine."
    return f"El proyecto {proyecto} no tiene cruces ({resumen})."


def _conectar_ro(bd: Path) -> sqlite3.Connection:
    con = sqlite3.connect(f"file:{bd}?mode=ro", uri=True)
    con.row_factory = sqlite3.Row
    return con


def cargar_nuestro(bd: Path, proyecto: int):
    """Cruces del proyecto por calzada y MINUTO, las motos aparte, y los
    minutos que el video realmente cubre: sin eso se compara contra horas no
    grabadas.

    Por minuto y no por cuarto de hora: los cuartos los pone el conteo de
    referencia. Si alguien conto leyendo la leyenda de la pantalla del aforo
    frontal, sus cuartos van corridos 4 h 55 min —que no es multiplo de 15—
    y caen en 13:05-13:20, no en 13:00-13:15. Con los cruces por minuto se
    arman los cuartos que haga falta.

    Devuelve (por_minuto, motos_por_minuto, cubiertos).
    """
    con = _conectar_ro(bd)

    # Solo los videos YA CONTADOS. Con la cola a medias, incluir los que
    # faltan hace creer que el video cubre tres horas cuando solo se han
    # contado veinte minutos, y entonces se comparan nuestros veinte
    # minutos contra tres horas de tubo: el resultado sale por los suelos
    # sin que nada este mal. Se acepta tambien un video sin cruces si ya
    # esta 'done', porque un tramo vacio de madrugada es un dato valido.
    cubiertos: set[int] = set()
    for r in con.execute(
        "select v.video_start_time t, v.total_frames f, v.fps from video_jobs v "
        "where v.project_id=? and v.video_start_time is not null "
        "  and (v.status = 'done' "
        "       or exists (select 1 from crossings x where x.job_id = v.id))",
        (proyecto,),
    ):
        h, m, _ = map(int, r["t"][11:].split(":"))
        # REDONDEADO, no truncado. Con segmentos de un minuto, un archivo de
        # 58 s (1 160 cuadros; hay 7 asi en el aforo frontal) daba int(0.97)
        # = 0 minutos: su minuto quedaba como "no grabado" y el cuarto de
        # hora entero salia de la comparacion sin avisar.
        dur = round((r["f"] or 0) / (r["fps"] or 15) / 60)
        cubiertos.update(range(h * 60 + m, h * 60 + m + dur))

    por_minuto: dict[str, dict[int, int]] = defaultdict(lambda: defaultdict(int))
    # El formato de la empresa no trae columna de motos (A, B, C, T-S,
    # T-S-R). Si no las cuentan, nuestro total lleva ~2 % de mas que ellos no
    # tienen; se llevan aparte para dar la razon con y sin ellas.
    motos: dict[str, dict[int, int]] = defaultdict(lambda: defaultdict(int))
    for zona, minuto, vt in _cruces_por_minuto(con, proyecto):
        por_minuto[zona][minuto] += 1
        if vt == "motorcycle":
            motos[zona][minuto] += 1
    con.close()
    return ({k: dict(v) for k, v in por_minuto.items()},
            {k: dict(v) for k, v in motos.items()}, cubiertos)


def agrupar(serie: dict[str, dict[int, int]], bins) -> dict[str, dict[int, int]]:
    """De cruces por minuto a cruces por cuarto de hora, con los cuartos que
    diga la referencia (que no tienen por que empezar en :00, :15, :30, :45)."""
    return {z: {b: sum(d.get(m % 1440, 0) for m in range(b, b + 15)) for b in bins}
            for z, d in serie.items()}


def _cruces_por_minuto(con, proyecto):
    """(zona, minuto del dia, clase de COCO) de cada cruce del proyecto."""
    for r in con.execute(
        "select z.name zona, x.timestamp ts, x.vehicle_type vt from crossings x "
        "join video_jobs v on v.id=x.job_id left join zones z on z.id=x.zone_id "
        "where v.project_id=?",
        (proyecto,),
    ):
        yield (r["zona"] or "sin zona",
               int(r["ts"][11:13]) * 60 + int(r["ts"][14:16]), r["vt"])


def fecha_del_proyecto(bd: Path, proyecto: int) -> str | None:
    con = _conectar_ro(bd)
    r = con.execute("select min(video_start_time) t from video_jobs "
                    "where project_id=? and video_start_time is not null",
                    (proyecto,)).fetchone()
    con.close()
    return r["t"][:10] if r and r["t"] else None


def buscar_desfase_reloj(por_minuto, real, pares, cubiertos, maximo=10):
    """Minutos que parece estar corrido el reloj del video respecto al conteo.

    La hora de cada cruce sale del nombre del archivo, que pone el reloj de
    la grabadora. En este material ese reloj ya dio una sorpresa: la leyenda
    impresa en la imagen se atraso casi un mes a las 12:25. Si el reloj esta
    corrido unos minutos, los vehiculos caen en el cuarto de hora vecino y la
    comparacion por intervalo empeora aunque el total cuadre.

    Se prueba correr nuestros cruces de -`maximo` a +`maximo` minutos y se
    mide el error por cuarto de hora contra el conteo. Es un DIAGNOSTICO: la
    cifra principal se da siempre sin correr nada, y si aparece un desfase lo
    que toca es preguntar a la empresa, no corregirlo a mano.

    Solo ve desfases de minutos. Uno de horas —como el de la leyenda del
    aforo frontal, 4 h 55 min— lo tiene que decir quien conto, y se aplica
    con --desfase-referencia; `revisar_reloj.py` dice cual reloj es el real.
    """
    candidatos = sorted({b for d in real.values() for b in d})
    combinado, por_zona = [], defaultdict(list)
    for s in range(-maximo, maximo + 1):
        # Solo cuartos de hora que sigan cubiertos enteros tras correrlos.
        bins = [b for b in candidatos
                if all(((m - s) % 1440) in cubiertos for m in range(b, b + 15))]
        error, base = 0.0, 0.0
        for zona, (sentido, _) in pares.items():
            if not sentido:
                continue
            ns = [sum(por_minuto.get(zona, {}).get((m - s) % 1440, 0)
                      for m in range(b, b + 15)) for b in bins]
            rs = [real[sentido].get(b, 0) for b in bins]
            # Se quita la escala antes de medir: si el sistema cuenta 0.9x,
            # el error absoluto nunca bajaria del 10 % y taparia el desfase.
            # Aqui solo importa si los vehiculos caen en su cuarto de hora.
            k = sum(rs) / sum(ns) if sum(ns) else 0.0
            e = sum(abs(k * n - r) for n, r in zip(ns, rs))
            if sum(rs):
                por_zona[zona].append((e / sum(rs), s))
            error += e
            base += sum(rs)
        if base:
            combinado.append((error / base, s))
    return sorted(combinado), {z: sorted(v) for z, v in por_zona.items()}


def veredicto_reloj(combinado, por_zona, mejora_minima=0.15):
    """Decide si hay un desfase de reloj de verdad, o si es ruido.

    Probar 21 desfases y quedarse con el mejor SIEMPRE le gana al cero por
    azar: con doce cuartos de hora, el minimo de 21 intentos casi nunca es
    el cero aunque el reloj este perfecto. La primera version cayo ahi y
    aviso "2 min adelantado" sobre el aforo viejo con una calzada rota.

    Lo que separa un desfase real del ruido es fisico: el reloj es UNO para
    todo el video, asi que si esta corrido, los dos sentidos por separado
    tienen que pedir el mismo desfase. Si cada uno pide uno distinto, lo que
    hay es ruido de cuarto de hora, no reloj.

    Devuelve (desfase o None, error en cero, error con desfase, texto).
    """
    if not combinado:
        return None, None, None, "sin datos"
    cero = next((e for e, s in combinado if s == 0), None)
    mejor_e, mejor_s = combinado[0]
    if cero is None or mejor_s == 0 or mejor_e >= (1 - mejora_minima) * cero:
        return None, cero, mejor_e, "sin desfase aparente"
    preferidos = {z: v[0][1] for z, v in por_zona.items() if v}
    if len(preferidos) >= 2 and all(abs(s - mejor_s) <= 1 for s in preferidos.values()):
        return mejor_s, cero, mejor_e, "los dos sentidos piden el mismo desfase"
    detalle = ", ".join(f"{z} {s:+d} min" for z, s in sorted(preferidos.items()))
    return None, cero, mejor_e, f"cada sentido pide otro desfase ({detalle}): ruido"


def _correlacion(a, b):
    """None cuando no se puede calcular: con pocos intervalos, o cuando una
    de las dos series es constante, la correlacion no esta definida."""
    if len(a) < 3:
        return None
    try:
        return statistics.correlation(a, b)
    except statistics.StatisticsError:
        return None


def _apenas_varia(serie, umbral: float = 0.12) -> bool:
    """Una serie casi constante no se puede correlacionar con nada de forma
    util: cualquier r que salga lo decide el ruido. El umbral es sobre el
    coeficiente de variacion (desviacion / media)."""
    if len(serie) < 2:
        return True
    media = statistics.fmean(serie)
    if media == 0:
        return True
    return statistics.pstdev(serie) / media < umbral


def emparejar(real, nuestro, bins):
    """Ata cada calzada nuestra al sentido real con el que su perfil temporal
    correlaciona mas.

    Se hace por los datos y no por el nombre porque "Calzada poniente" es
    ambiguo: puede ser la calzada del lado poniente o la que lleva al
    poniente, y equivocarse invierte la comparacion entera.

    El emparejamiento es UNO A UNO. Antes cada calzada elegia su mejor
    sentido por separado, y con pocos intervalos —cuando la correlacion no
    se puede calcular y todas empatan— las dos calzadas se quedaban con el
    mismo sentido: el informe mostraba el mismo "real" dos veces y los
    porcentajes no sumaban 100.
    """
    puntajes = []
    for zona, datos in nuestro.items():
        v = [datos.get(b, 0) for b in bins]
        for sentido, d in real.items():
            c = _correlacion([d.get(b, 0) for b in bins], v)
            puntajes.append((c if c is not None else -2.0, c, zona, sentido))

    pares, zonas_por_atar, sentidos_libres = {}, set(nuestro), set(real)
    for _, c, zona, sentido in sorted(puntajes, key=lambda p: -p[0]):
        if zona in zonas_por_atar and sentido in sentidos_libres:
            pares[zona] = (sentido, c)
            zonas_por_atar.discard(zona)
            sentidos_libres.discard(sentido)
    # Si sobran calzadas (mas calzadas que sentidos) quedan sin pareja.
    for zona in zonas_por_atar:
        pares[zona] = (None, None)
    return pares


def composicion_nuestra(bd: Path, proyecto: int, bins) -> dict[str, int]:
    """Clases de la ventana comparable, TAL COMO LAS VE EL USUARIO.

    Se leen de `traffic_db.get_interval_counts`, que es el unico sitio donde
    las clases de COCO se traducen a la taxonomia de la empresa (pantallas,
    graficas y los dos CSV pasan por ahi). Comparar otra cosa seria validar
    una cifra que nadie ve.
    """
    sys.path.insert(0, str(RAIZ))
    from src.storage import traffic_db
    traffic_db.DB_PATH = Path(bd)
    datos = traffic_db.get_interval_counts(proyecto, 15)
    comp: dict[str, int] = defaultdict(int)
    for carril in datos["lanes"]:
        for it in carril["intervals"]:
            t = it["start"]
            # Los intervalos de la plataforma van en :00/:15/:30/:45; los de
            # la referencia pueden ir corridos (un conteo hecho con la hora de
            # la pantalla). Entra el intervalo cuyo punto medio cae dentro de
            # un cuarto comparable: para una PROPORCION, cinco minutos de
            # orilla no pesan.
            medio = int(t[11:13]) * 60 + int(t[14:16]) + 7.5
            if not any(b <= medio < b + 15 for b in bins):
                continue
            for clase, v in (it.get("by_vehicle_type") or {}).items():
                comp[clase] += v.get("in", 0) + v.get("out", 0)
    return dict(comp)


def informe_clases(clases_real, pares, bins, comp) -> None:
    """Composicion por clase contra el conteo manual, los dos sentidos juntos.

    El conteo manual trae A, B, C, T-S y T-S-R. La camara frontal entrega
    MOTO / A / B / C: su C son los camiones de un cuerpo Y los articulados,
    que desde la camara no se separan. La camara vieja entrega A / PESADO.
    """
    real = defaultdict(float)
    for sentido, _ in pares.values():
        if not sentido or sentido not in clases_real:
            continue
        for b in bins:
            for clase, v in clases_real[sentido].get(b, {}).items():
                real[clase] += v
    sin_resolver = comp.get("SIN_RESOLVER", 0)
    if "PESADO" in comp:
        # Con la camara de clases gruesas nuestras motos van dentro de A;
        # las del contador de ejes tambien, o se comparan cosas distintas.
        real["A"] += real.pop("MOTO", 0.0)
        filas = [("A (livianos)", real["A"], comp.get("A", 0)),
                 ("Pesados", real["B"] + real["C"] + real["T-S"] + real["T-S-R"],
                  comp.get("PESADO", 0))]
    else:
        # Del lado nuestro el C se junta con T-S y el tractor sin caja: la
        # empresa no dice en que columna cuenta al tractor, asi que separarlos
        # compararia repartos distintos de los mismos vehiculos.
        filas = [("A (livianos)", real["A"], comp.get("A", 0)),
                 ("B autobus", real["B"], comp.get("B", 0)),
                 ("C camion (C+T-S+T-S-R+tractor)",
                  real["C"] + real["T-S"] + real["T-S-R"],
                  sum(comp.get(k, 0) for k in ("C", "T-S", "T-S-R", "TRACTOR")))]
    motos = comp.get("MOTO", 0)
    # El contador de ejes trae motos ("Cycle"): entonces entran a la tabla
    # como una clase mas, de los dos lados.
    con_motos = "MOTO" in real
    if con_motos:
        filas.append(("Motocicleta", real["MOTO"], motos))
        motos = 0
    tr = sum(f[1] for f in filas)
    tn = sum(f[2] for f in filas)
    if not tr or not tn:
        print("\nComposicion: sin datos de clase en la ventana.")
        return
    titulo = "con motos" if con_motos else "sin motos"
    print(f"\n=== Composicion por clase (ambos sentidos, {titulo}) ===")
    print(f"{'clase':<26}{'nuestro':>9}{'real':>9}{'razon':>8}"
          f"{'% nuestro':>11}{'% real':>9}{'puntos':>8}")
    for nombre, r, n in filas:
        razon = f"{n / r:.2f}x" if r else "-"
        pn, pr = 100 * n / tn, 100 * r / tr
        print(f"{nombre:<26}{n:>9}{int(r):>9}{razon:>8}{pn:>10.1f}%"
              f"{pr:>8.1f}%{abs(pn - pr):>8.1f}")
    if motos:
        print(f"Motocicletas nuestras: {motos}. El formato de la empresa no las "
              "trae; si las\n  cuentan dentro de A, sumarlas a nuestra A antes "
              "de comparar.")
    if real.get("OTRO"):
        print(f"El contador de ejes dejo {int(real['OTRO'])} vehiculos como 'Other' "
              "(ejes que no forman un patron conocido), fuera de los porcentajes.")
    if sin_resolver:
        print(f"Sin clasificar (noche o vehiculo muy chico): {sin_resolver}, "
              "fuera de los porcentajes.")
    print("  (La composicion puede cuadrar con errores que se cancelan: una "
          "troca contada\n   como pesada y un camion como liviano. Para eso "
          "estan las hojas de recortes.)")


def main() -> None:
    p = argparse.ArgumentParser(
        description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter
    )
    p.add_argument("--proyecto", type=int, required=True)
    p.add_argument("--referencias", type=Path, default=REFERENCIAS)
    p.add_argument("--fuente", choices=("auto", "manual", "tubo"), default="auto",
                   help="referencia a usar; auto prefiere el conteo manual")
    p.add_argument("--bd", type=Path, default=BD)
    p.add_argument("--salida", type=Path, help="CSV con el detalle por cuarto de hora")
    p.add_argument("--ignorar-fecha", action="store_true",
                   help="comparar aunque la fecha del conteo no sea la del video")
    p.add_argument("--desfase-referencia", type=int, default=0, metavar="MIN",
                   help="minutos que hay que SUMAR a las horas del conteo para "
                        "llevarlas a la hora real. Aforo frontal del 19-sep-2026 "
                        "contado con la leyenda de la pantalla: -295 (4 h 55 min)")
    a = p.parse_args()

    print("Aforo real de referencia:")
    fecha_video = fecha_del_proyecto(a.bd, a.proyecto)
    real, clases_real, fecha_ref = cargar_referencia(a.referencias, a.fuente, fecha_video)
    if a.desfase_referencia:
        # Se corre la REFERENCIA, no lo nuestro: nuestras horas salen del
        # nombre del archivo, que `revisar_reloj.py` comprobo contra el sol.
        d = a.desfase_referencia
        real = {s: {(m + d) % 1440: v for m, v in x.items()} for s, x in real.items()}
        if clases_real:
            clases_real = {s: {(m + d) % 1440: v for m, v in x.items()}
                           for s, x in clases_real.items()}
        print(f"  horas del conteo corridas {d:+d} min para llevarlas a la hora real")
        # La fecha del titulo es la que se leyo en la pantalla; con la
        # leyenda del frontal tambien viene corrida (22 dias), asi que ya no
        # sirve de candado.
        fecha_ref = None
    # La carpeta por omision trae el aforo del 19-ago-2026 (camara lateral).
    # Sin este candado, el proyecto frontal del 19-sep se compara contra el
    # dia equivocado y la cifra sale razonable y falsa a la vez.
    if fecha_ref and fecha_video and fecha_ref != fecha_video and not a.ignorar_fecha:
        print()
        sys.exit(f"El conteo es del {fecha_ref} y el video del {fecha_video}. "
                 "No se comparan dias distintos.\nPon el conteo de este video "
                 "en su propia carpeta (--referencias referencias/<aforo>/),\n"
                 "o usa --ignorar-fecha si de verdad es el mismo dia.")
    por_minuto, motos_minuto, cubiertos = cargar_nuestro(a.bd, a.proyecto)
    if not por_minuto:
        sys.exit(diagnosticar_sin_cruces(a.bd, a.proyecto))

    # Solo los cuartos que la referencia CONTO y que el video cubre ENTEROS:
    # uno a medias compara 15 minutos de referencia contra 2 de camara.
    bins = sorted(
        b for b in {m for x in real.values() for m in x}
        if all((m % 1440) in cubiertos for m in range(b, b + 15))
    )
    if not bins:
        sys.exit("El video no cubre ningun cuarto de hora de los que trae el conteo.")
    nuestro = agrupar(por_minuto, bins)
    print(
        f"\nVentana comparable: {bins[0] // 60:02d}:{bins[0] % 60:02d} a "
        f"{(bins[-1] + 15) // 60:02d}:{(bins[-1] + 15) % 60:02d}  "
        f"({len(bins)} cuartos de hora)"
    )

    pares = emparejar(real, nuestro, bins)
    print("\nCalzada nuestra -> sentido real (emparejado por correlacion del perfil):")
    for zona, (sentido, r) in sorted(pares.items()):
        marca = f"r = {r:+.2f}" if r is not None else "r no calculable"
        print(f"  {zona:<20} -> {str(sentido):<12} {marca}")
    if len(bins) < 4:
        print(f"  AVISO: solo {len(bins)} cuartos de hora contados. Con tan pocos, el")
        print("  emparejamiento calzada-sentido no es fiable; puede estar invertido.")

    print("\n=== Totales en la ventana comparable ===")
    # Las motos se quitan tambien POR CALZADA, no solo del total. La prueba
    # con un conteo sin motos lo destapo: cada sentido salia 3 % inflado
    # (0.96x donde se habia metido 0.93x) aunque el total "sin motos" cuadraba.
    motos = agrupar(motos_minuto, bins)
    # El conteo manual no trae motos; el contador de ejes SI (su clase
    # "Cycle"). Con el tubo, "sin motos" las quita de los DOS lados: es la
    # comparacion sin la clase mas dificil para ambos instrumentos.
    ref_con_motos = bool(clases_real) and any(
        "MOTO" in v for x in clases_real.values() for v in x.values())

    def motos_ref(sentido, b):
        if not ref_con_motos or not sentido:
            return 0
        return clases_real.get(sentido, {}).get(b, {}).get("MOTO", 0)

    print(f"{'calzada':<24}{'nuestro':>9}{'real':>9}{'razon':>8}{'sin motos':>11}")
    tn = tr = mn = mr = 0
    for zona, (sentido, _) in sorted(pares.items()):
        n = sum(nuestro[zona].get(b, 0) for b in bins)
        rr = sum(real[sentido].get(b, 0) for b in bins) if sentido else 0
        mz = sum(motos.get(zona, {}).get(b, 0) for b in bins)
        rz = sum(motos_ref(sentido, b) for b in bins)
        tn, tr, mn, mr = tn + n, tr + rr, mn + mz, mr + rz
        razon = f"{n / rr:.2f}x" if rr else "-"
        sin = f"{(n - mz) / (rr - rz):.2f}x" if rr - rz else "-"
        print(f"{zona:<24}{n:>9}{int(rr):>9}{razon:>8}{sin:>11}")
    if tr:
        print(f"{'AMBOS SENTIDOS':<24}{tn:>9}{int(tr):>9}{tn / tr:>7.2f}x"
              f"{(tn - mn) / (tr - mr):>10.2f}x")
        if ref_con_motos:
            print(f"  (motocicletas: {mn} nuestras, {int(mr)} del contador de ejes.) El "
                  "contador SI las cuenta:")
            print("  la razon buena es la primera; la de 'sin motos' compara sin "
                  "ellas de los dos lados.")
        elif mn:
            print(f"  ({mn} motocicletas nuestras.) El conteo de la empresa no "
                  "trae columna de motos:")
            print("  si NO las cuentan, la razon buena es la de 'sin motos'; si "
                  "las meten en A,")
            print("  la otra. Hay que preguntarlo.")

    print("\n=== Reparto entre sentidos ===")
    reparto_real = " / ".join(
        f"{100 * sum(real[s].get(b, 0) for b in bins) / tr:.0f}% {s}" for s in sorted(real)
    )
    reparto_nuestro = " / ".join(
        f"{100 * sum(nuestro[z].get(b, 0) for b in bins) / tn:.0f}% {z}"
        for z in sorted(nuestro)
    )
    print(f"  real:    {reparto_real}")
    print(f"  nuestro: {reparto_nuestro}")

    sr = [sum(d.get(b, 0) for d in real.values()) for b in bins]
    sn = [sum(d.get(b, 0) for d in nuestro.values()) for b in bins]
    r_total = _correlacion(sr, sn)
    if r_total is None:
        print("\nCorrelacion del perfil temporal: no se puede calcular todavia.")
        print(f"  Hacen falta al menos 3 cuartos de hora contados; hay {len(bins)}.")
    elif _apenas_varia(sr):
        # Sumar los dos sentidos puede aplanar el perfil: uno sube mientras
        # el otro baja y el total queda casi constante. Correlacionar contra
        # una serie plana da un numero que parece informativo y no lo es —
        # se midio un r = -0.55 sobre cuatro intervalos donde el total real
        # iba 373, 373, 400, 394, mientras cada sentido por separado
        # correlacionaba a +0.97 y +0.87.
        print(f"\nCorrelacion del perfil temporal (ambos sentidos): r = {r_total:+.2f}")
        print("  NO INTERPRETABLE: el total real apenas varia entre intervalos")
        print(f"  ({min(sr):.0f} a {max(sr):.0f}), asi que este r es ruido. Mira arriba")
        print("  la correlacion de cada sentido por separado, que si dice algo.")
    else:
        print(f"\nCorrelacion del perfil temporal (ambos sentidos): r = {r_total:+.2f}")
        print("  r alto con razon lejos de 1 = el detector si ve el transito real,")
        print("  pero la escala esta mal. r bajo = no lo esta viendo.")

    # El reloj de la grabadora puede estar corrido. Se reporta, no se corrige.
    combinado, por_zona = buscar_desfase_reloj(por_minuto, real, pares, cubiertos)
    desfase, cero, mejor, motivo = veredicto_reloj(combinado, por_zona)
    print()
    if desfase is not None:
        lado = "adelantado" if desfase < 0 else "atrasado"
        print(f"AVISO DE RELOJ: el reloj del video parece ir {abs(desfase)} min {lado}")
        print(f"  ({motivo}; el error por cuarto de hora baja de "
              f"{100 * cero:.0f} % a {100 * mejor:.0f} %).")
        print("  Las cifras de arriba NO lo corrigen: preguntar a la empresa si "
              "la grabadora")
        print("  tenia la hora sincronizada.")
    elif cero is not None:
        print(f"Reloj: {motivo} (error por cuarto de hora {100 * cero:.0f} % "
              f"sin correr nada).")

    if clases_real:
        informe_clases(clases_real, pares, bins,
                       composicion_nuestra(a.bd, a.proyecto, bins))

    if a.salida:
        a.salida.parent.mkdir(parents=True, exist_ok=True)
        with open(a.salida, "w", newline="", encoding="utf-8") as f:
            w = csv.writer(f)
            zonas = sorted(nuestro)
            w.writerow(
                ["hora"]
                + [f"real_{pares[z][0]}" for z in zonas]
                + [f"nuestro_{z}" for z in zonas]
            )
            for b in bins:
                w.writerow(
                    [f"{b // 60:02d}:{b % 60:02d}"]
                    + [int(real[pares[z][0]].get(b, 0)) for z in zonas]
                    + [nuestro[z].get(b, 0) for z in zonas]
                )
        print(f"\nDetalle por cuarto de hora en {a.salida}")


if __name__ == "__main__":
    main()
