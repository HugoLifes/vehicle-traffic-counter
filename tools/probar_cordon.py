"""
Regresión sin GPU del cordón de líneas (traffic_db.es_cordon).

En un aforo direccional con una línea suelta por brazo, cada vehículo cruza
DOS líneas: la de su entrada y la de su salida. El volumen del cruce y su hora
de máxima demanda tienen que salir de las entradas, en el Reporte
(traffic_metrics) y en el Excel (hoja RESUMEN) por igual; sumando todos los
cruces salían al doble.

Fabrica un cruce de tres brazos en una base temporal con 400 vehículos, 100
de ellos amontonados en 14:00-14:15, y comprueba:
  1. que el proyecto se reconozca como cordón, y que deje de serlo al atar
     las líneas a calzadas (el aforo por línea de siempre no cambia);
  2. volumen = 400, no 800;
  3. la misma hora de máxima demanda en el Reporte y en el Excel, y que sea
     la que contiene el amontonamiento.

    python tools/probar_cordon.py
"""

import random
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from src.storage import traffic_db  # noqa: E402

fallos = 0


def comprobar(nombre, ok, detalle=""):
    global fallos
    fallos += 0 if ok else 1
    print(f"  [{'ok' if ok else 'FALLA'}] {nombre}" + (f" — {detalle}" if detalle else ""))


def main():
    tmp = Path(tempfile.mkdtemp())
    traffic_db.DB_PATH = tmp / "cordon.db"
    traffic_db.init_schema()
    pid = traffic_db.create_project("Cordón de prueba", interval_minutes=15)
    for n in "ABC":
        traffic_db.create_zone(pid, n, [[0, 0], [1, 0], [1, 1]], kind="acceso")
    lineas = {n: traffic_db.create_lane("x", n, "line", [[0, 0], [0, 1]], project_id=pid)
              for n in "ABC"}
    jid = traffic_db.create_video_job("v.mp4", "v.mp4", project_id=pid,
                                      video_start_time="2026-10-07 13:00:00")
    traffic_db.update_video_job(jid, fps=15, total_frames=15 * 3600 * 2, status="done")
    azar = random.Random(1)
    for i in range(400):
        minuto = azar.randint(0, 119) if i < 300 else azar.randint(60, 74)
        t = f"2026-10-07 {13 + minuto // 60:02d}:{minuto % 60:02d}:{azar.randint(0, 59):02d}"
        a, b = azar.sample("ABC", 2)
        traffic_db.record_crossing(lineas[a], i, "in", "car", 0.8, timestamp=t, job_id=jid)
        traffic_db.record_crossing(lineas[b], i, "out", "car", 0.8, timestamp=t, job_id=jid)

    comprobar("se reconoce como cordón", traffic_db.es_cordon(pid))

    from src.storage import traffic_metrics
    m = traffic_metrics.get_project_metrics(pid, 15)
    comprobar("volumen del cruce = vehículos", m["totals"]["total"] == 400,
              f"{m['totals']['total']} (cruces de línea: {m['totals'].get('cruces')})")
    pico = m["peak_hour"]
    comprobar("la hora pico contiene el amontonamiento de 14:00",
              pico and pico["start"] <= "2026-10-07 14:00:00" < pico["end"], pico and pico["start"])

    from openpyxl import load_workbook
    from src.reports import aforo_excel
    ruta = tmp / "cordon.xlsx"
    aforo_excel.generar(pid, str(ruta))
    filas = [f for f in load_workbook(ruta)["RESUMEN"].iter_rows(values_only=True)
             if f[1] and str(f[1]).startswith("TODO EL CRUCE")]
    comprobar("el Excel trae la fila del cruce", len(filas) == 1)
    if filas:
        _, _, total, hmd, volumen, _ = filas[0][:6]
        comprobar("Excel: volumen = vehículos", total == 400, str(total))
        comprobar("Excel y Reporte dan la misma hora pico",
                  hmd.startswith(pico["start"][11:16]) and volumen == pico["volume"],
                  f"Excel {hmd} {volumen}, Reporte {pico['start'][11:16]} {pico['volume']}")

    # Atadas a calzadas ya no es un cordón: el aforo por línea de siempre.
    calzada = traffic_db.create_zone(pid, "Calzada", [[0, 0], [1, 0], [1, 1]], kind="calzada")
    for lid in lineas.values():
        traffic_db.get_connection().execute(
            "UPDATE lane_configs SET zone_id = ? WHERE id = ?", (calzada, lid))
    traffic_db.get_connection().commit()
    comprobar("con calzadas deja de ser cordón", not traffic_db.es_cordon(pid))

    print(f"\n{fallos} fallos")
    sys.exit(1 if fallos else 0)


if __name__ == "__main__":
    main()
