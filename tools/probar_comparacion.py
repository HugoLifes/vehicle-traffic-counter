"""
Prueba de punta a punta de `comparar_aforo_real.py`, sobre un proyecto real
y ANTES de que llegue el conteo manual de la empresa.

Fabrica conteos manuales con el formato exacto que manda la empresa (dos
hojas, "Hr/Mov", bloques AM y PM, clases A B C T-S T-S-R, la fecha en el
titulo) a partir de los cruces del propio proyecto, con errores CONOCIDOS
metidos a proposito, y comprueba que la herramienta los encuentra. Tres
escenarios, cada uno sacado de algo que de verdad puede pasar:

1. Conteo del dia completo con el reloj corrido 3 min, un sentido al 0.93x
   y el otro al 0.98x, sin motos y con sentidos escritos NTE-SUR.
2. Conteo de SOLO TRES FRANJAS, el resto de la hoja en blanco. Es lo que se
   le pidio a la empresa. Una fila en blanco no es "cero vehiculos": la
   herramienta la leia como cero y comparaba trece horas nuestras contra
   ceros.
3. Conteo hecho LEYENDO LA PANTALLA. En el aforo frontal la leyenda salto a
   las 12:25 (22 dias atras, 4 h 55 min adelante): quien cuente con ella
   pone sus cuartos de hora corridos y con fecha del 28 de agosto. Sin
   decirle nada, la herramienta tiene que negarse (la fecha no cuadra); con
   --desfase-referencia -295 tiene que recuperar las razones, aunque 4 h 55
   no sea multiplo de 15 y sus cuartos caigan en 13:05-13:20.

Y que se niega a comparar contra el conteo de otro dia (el del 19-ago).

    python tools/probar_comparacion.py --proyecto 7 --salida data/prueba_comparacion

Hasta ahora esta prueba encontro: 7 videos de 58 s que tiraban sus cuartos
de hora, el aviso de reloj gritando lobo sobre un aforo sin desfase, las
motos inflando 3 % cada sentido, y las filas en blanco leidas como cero.
"""

import argparse
import re
import shutil
import sqlite3
import statistics
import subprocess
import sys
from collections import defaultdict
from pathlib import Path

RAIZ = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(RAIZ))

RAZON = (0.93, 0.98)        # lo que el sistema "captura" de cada sentido
SENTIDOS = ("NTE-SUR", "SUR-NTE")
LEYENDA = 295               # 4 h 55 min: lo que adelanta la leyenda desde las 12:25
MESES = ("ENERO", "FEBRERO", "MARZO", "ABRIL", "MAYO", "JUNIO", "JULIO",
         "AGOSTO", "SEPTIEMBRE", "OCTUBRE", "NOVIEMBRE", "DICIEMBRE")


def etiqueta(minuto):
    def hhmm(m):
        h, mi = divmod(m % 1440, 60)
        return f"{h % 12 or 12}:{mi:02d} {'AM' if h < 12 else 'PM'}"
    return f"{hhmm(minuto)} - {hhmm(minuto + 15)}"


def leer_proyecto(bd, proyecto):
    con = sqlite3.connect(f"file:{bd}?mode=ro", uri=True)
    con.row_factory = sqlite3.Row
    zonas = [r["name"] for r in con.execute(
        "select distinct z.name from crossings x join zones z on z.id=x.zone_id "
        "join video_jobs v on v.id=x.job_id where v.project_id=? order by 1",
        (proyecto,))]
    if len(zonas) != 2:
        sys.exit(f"La prueba necesita un proyecto de dos calzadas; este tiene {zonas}")
    umbral = {}
    for z in zonas:
        altos = [r[0] for r in con.execute(
            "select x.bbox_height from crossings x join zones z on z.id=x.zone_id "
            "where z.name=? and x.vehicle_type='car' and x.bbox_height is not null", (z,))]
        umbral[z] = 1.58 * statistics.median(altos)
    cruces = [dict(r) for r in con.execute(
        "select z.name zona, x.timestamp ts, x.vehicle_type vt, x.bbox_height h "
        "from crossings x join zones z on z.id=x.zone_id "
        "join video_jobs v on v.id=x.job_id where v.project_id=?", (proyecto,))]
    # Horas cubiertas enteras por videos ya contados: de ahi salen las franjas.
    minutos = set()
    fecha = None
    for r in con.execute("select video_start_time t, total_frames f, fps from video_jobs "
                         "where project_id=? and status='done'", (proyecto,)):
        fecha = fecha or r["t"][:10]
        h, m = int(r["t"][11:13]), int(r["t"][14:16])
        minutos.update(range(h * 60 + m, h * 60 + m + round(r["f"] / r["fps"] / 60)))
    con.close()
    horas = [h for h in range(24) if all(h * 60 + m in minutos for m in range(60))]
    return zonas, umbral, cruces, horas, fecha


def fabricar(cruces, zonas, umbral, fecha, ruta, reloj=0, leyenda=0, ventanas=None):
    """Escribe un conteo manual de prueba.

    reloj: minutos que va adelantada la grabadora (el conteo lleva la hora real).
    leyenda: minutos que adelanta la PANTALLA con la que se conto (sus
        etiquetas = hora real + leyenda, y sus cuartos se arman sobre ese reloj).
    ventanas: [(desde, hasta)] en minutos DE LA ETIQUETA; fuera, filas en blanco.
    """
    import openpyxl
    conteo = {s: defaultdict(lambda: defaultdict(float)) for s in SENTIDOS}
    alterna = 0
    for r in cruces:
        if r["vt"] == "motorcycle":
            continue                       # la empresa no las cuenta
        i = zonas.index(r["zona"])
        real = int(r["ts"][11:13]) * 60 + int(r["ts"][14:16]) - reloj
        marca = (real + leyenda) % 1440
        if r["vt"] == "bus":
            clase = "B"
        elif r["vt"] == "truck" and (r["h"] or 0) > umbral[r["zona"]]:
            alterna += 1
            clase = "C" if alterna % 2 else "T-S"
        else:
            clase = "A"
        conteo[SENTIDOS[i]][(marca // 15) * 15][clase] += 1 / RAZON[i]

    dia = fecha
    if leyenda:
        # La pantalla del frontal tambien mentia en la fecha: 22 dias atras.
        import datetime as dt
        dia = (dt.date.fromisoformat(fecha) - dt.timedelta(days=22)).isoformat()
    titulo = f"CRUCE DE PRUEBA {int(dia[8:])}-{MESES[int(dia[5:7]) - 1]}-{dia[:4]}"

    libro = openpyxl.Workbook()
    libro.remove(libro.active)
    for s in SENTIDOS:
        h = libro.create_sheet(s)
        for col0 in (2, 9):               # bloque AM en B, bloque PM en I
            h.cell(3, col0, titulo)
            h.cell(4, col0, "Hr/Mov")
            h.cell(4, col0 + 1, s)
            for i, c in enumerate(("A", "B", "C", "T-S", "T-S-R"), start=1):
                h.cell(5, col0 + i, c)
        for k in range(48):
            for col0, base in ((2, 0), (9, 720)):
                m = base + 15 * k
                h.cell(6 + k, col0, etiqueta(m))
                if ventanas and not any(d <= m < ht for d, ht in ventanas):
                    continue              # nadie conto este cuarto: en blanco
                for i, c in enumerate(("A", "B", "C", "T-S", "T-S-R"), start=1):
                    h.cell(6 + k, col0 + i, round(conteo[s][m][c]))
    shutil.rmtree(ruta.parent, ignore_errors=True)
    ruta.parent.mkdir(parents=True, exist_ok=True)
    libro.save(ruta)


def correr(args):
    r = subprocess.run([sys.executable, str(RAIZ / "tools" / "comparar_aforo_real.py")]
                       + args, capture_output=True, text=True)
    return r.returncode, r.stdout + r.stderr


def main():
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--proyecto", type=int, required=True)
    ap.add_argument("--bd", default=str(RAIZ / "data" / "traffic.db"))
    ap.add_argument("--salida", default=str(RAIZ / "data" / "prueba_comparacion"))
    ap.add_argument("--ver", action="store_true", help="imprimir la salida completa")
    a = ap.parse_args()

    zonas, umbral, cruces, horas, fecha = leer_proyecto(a.bd, a.proyecto)
    # Tres franjas repartidas, con la hora SIGUIENTE tambien contada (los
    # cuartos de la hora de pantalla se salen 5 min a ella) y sin que la
    # etiqueta de pantalla pase de medianoche (la hoja es de un solo dia).
    utiles = [h for h in horas if h + 1 in horas and h * 60 + 60 + LEYENDA <= 1440]
    if len(utiles) < 3:
        sys.exit("Hacen falta mas horas contadas enteras para la prueba")
    tres = [utiles[0], utiles[len(utiles) // 2], utiles[-1]]
    base = Path(a.salida)
    fallos = 0

    def comprobar(ok, que):
        nonlocal fallos
        print(("ok   " if ok else "MAL  ") + que)
        fallos += 0 if ok else 1

    def razones(texto):
        out = {}
        for z in zonas:
            m = re.search(rf"{re.escape(z)}\s+\d+\s+\d+\s+[\d.]+x\s+([\d.]+)x", texto)
            out[z] = float(m.group(1)) if m else None
        return out

    def cuartos(texto):
        m = re.search(r"\((\d+) cuartos de hora\)", texto)
        return int(m.group(1)) if m else None

    def revisar_razones(texto, etiqueta_):
        for z, razon in zip(zonas, RAZON):
            got = razones(texto)[z]
            comprobar(got is not None and abs(got - razon) <= 0.015,
                      f"{etiqueta_}: razon sin motos de {z} {got}x (se metio {razon}x)")

    # --- 1. Dia completo, reloj corrido 3 min ------------------------------
    print("\n== 1. Dia completo, reloj 3 min adelantado ==")
    d1 = base / "dia_completo"
    fabricar(cruces, zonas, umbral, fecha, d1 / "conteo.xlsx", reloj=3)
    codigo, texto = correr(["--proyecto", str(a.proyecto), "--bd", a.bd,
                            "--referencias", str(d1)])
    if a.ver:
        print(texto)
    comprobar(codigo == 0, "corre sobre el proyecto")
    comprobar("nte-sur" in texto and "sur-nte" in texto, "reconoce sentidos NTE-SUR")
    comprobar("reloj del video parece ir 3 min adelantado" in texto,
              "encuentra el reloj corrido 3 min")
    revisar_razones(texto, "dia completo")
    comprobar("Composicion por clase" in texto and "B autobus" in texto,
              "compara MOTO / A / B / C contra A B C T-S T-S-R")

    # --- 2. Solo tres franjas ----------------------------------------------
    print(f"\n== 2. Solo tres franjas: {', '.join(f'{h:02d}:00' for h in tres)} ==")
    d2 = base / "tres_franjas"
    fabricar(cruces, zonas, umbral, fecha, d2 / "conteo.xlsx",
             ventanas=[(h * 60, h * 60 + 60) for h in tres])
    codigo, texto = correr(["--proyecto", str(a.proyecto), "--bd", a.bd,
                            "--referencias", str(d2)])
    if a.ver:
        print(texto)
    comprobar(codigo == 0, "corre con la hoja casi toda en blanco")
    comprobar(cuartos(texto) == 12,
              f"compara solo los 12 cuartos contados (dice {cuartos(texto)})")
    revisar_razones(texto, "tres franjas")

    # --- 3. Contado leyendo la pantalla ------------------------------------
    print(f"\n== 3. Contado con la hora de la pantalla (+{LEYENDA} min, fecha -22 dias) ==")
    d3 = base / "hora_de_pantalla"
    fabricar(cruces, zonas, umbral, fecha, d3 / "conteo.xlsx", leyenda=LEYENDA,
             ventanas=[(h * 60 + LEYENDA, h * 60 + 60 + LEYENDA) for h in tres])
    codigo, texto = correr(["--proyecto", str(a.proyecto), "--bd", a.bd,
                            "--referencias", str(d3)])
    comprobar(codigo != 0 and "No se comparan dias distintos" in texto,
              "sin avisarle, se niega: la fecha de la pantalla no es la del video")
    codigo, texto = correr(["--proyecto", str(a.proyecto), "--bd", a.bd,
                            "--referencias", str(d3), "--desfase-referencia", str(-LEYENDA)])
    if a.ver:
        print(texto)
    comprobar(codigo == 0, "con --desfase-referencia -295 corre")
    comprobar(cuartos(texto) == 12,
              f"arma los 12 cuartos corridos 5 min (dice {cuartos(texto)})")
    revisar_razones(texto, "hora de pantalla")

    # --- 4. Contra el conteo de otro dia -----------------------------------
    print("\n== 4. Contra el conteo del 19-ago ==")
    codigo, texto = correr(["--proyecto", str(a.proyecto), "--bd", a.bd])
    comprobar(codigo != 0 and "No se comparan dias distintos" in texto,
              "se niega a comparar contra el conteo de otro dia")

    print(f"\n{fallos} fallos")
    return 1 if fallos else 0


if __name__ == "__main__":
    sys.exit(main())
