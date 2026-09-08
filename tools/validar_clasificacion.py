#!/usr/bin/env python3
"""
Valida la clasificacion vehicular contra el aforo contado a mano.

El problema que resuelve: la clase `truck` de COCO, que es con la que sale
YOLO, mete en el mismo saco una pickup y un tractocamion. En la
clasificacion que usa la empresa (SCT: A, B, C, T-S, T-S-R) la pickup es
"A" (automovil) y el tractocamion es "T-S". Comparadas en crudo, nuestro
32 % de `truck` contra el 12.6 % de pesados del conteo manual parece un
error garrafal, y no lo es: son dos taxonomias distintas.

La separacion se hace por ALTO EN PIXELES, y funciona por una razon
concreta de este montaje: todos los vehiculos se cuentan al cruzar una
linea fija, o sea practicamente a la misma distancia de la camara. A
distancia constante, el alto en pixeles es proporcional al alto real.

El umbral no se elige a ojo: se calibra en la calzada cercana contra el
conteo manual y se expresa como multiplo del alto MEDIANO del automovil de
esa calzada, para que valga tambien donde la escala es otra. La calzada
del fondo sirve de validacion independiente: su umbral sale de la regla,
no de sus propios datos.

Uso:
    python tools/validar_clasificacion.py --proyecto 2
"""
from __future__ import annotations

import argparse
import re
import sqlite3
import statistics
import sys
from pathlib import Path

RAIZ = Path(__file__).resolve().parent.parent
BD = RAIZ / "data" / "traffic.db"
MANUAL = RAIZ / "referencias" / "aforo_real" / "conteo_manual_24h.xlsx"

# Clases del conteo manual que son vehiculo pesado. "A" (automoviles)
# incluye pickups y camionetas, que es justo lo que YOLO llama `truck`.
PESADAS = ("B", "C", "T-S", "T-S-R")

# YOLO no distingue autobus de camion de forma fiable a este tamano: en la
# calzada cercana los `bus` miden 65-90 px y los `truck` grandes tambien.
# Por eso se valida el corte livianos/pesados y no la clase fina.
LIVIANAS_YOLO = ("car", "motorcycle")


def leer_manual(ruta: Path, bins: list[int]) -> dict[str, dict[str, int]]:
    """{sentido: {clase: total en la ventana}}"""
    import openpyxl

    libro = openpyxl.load_workbook(str(ruta), data_only=True)
    fuera: dict[str, dict[str, int]] = {}
    for hoja in libro.worksheets:
        cols, sentido = [], None
        for fila in hoja.iter_rows(min_row=1, max_row=6):
            for celda in fila:
                txt = str(celda.value or "").strip()
                if txt.lower().startswith("hr/mov"):
                    cols.append(celda.column)
                elif re.fullmatch(r"(OTE|PTE)-(OTE|PTE)", txt, re.I):
                    sentido = txt.lower()
        if not cols or not sentido:
            continue
        acum: dict[str, int] = {}
        vistos: set[int] = set()
        for fila in range(1, hoja.max_row + 1):
            for col in cols:
                m = re.match(r"\s*(\d{1,2}):(\d{2})\s*(AM|PM)",
                             str(hoja.cell(fila, col).value or ""), re.I)
                if not m:
                    continue
                h, mi, ap = int(m.group(1)), int(m.group(2)), m.group(3).upper()
                if ap == "PM" and h != 12:
                    h += 12
                if ap == "AM" and h == 12:
                    h = 0
                minuto = h * 60 + mi
                if minuto in vistos or minuto not in bins:
                    continue
                vistos.add(minuto)
                for i, clase in enumerate(("A",) + PESADAS, start=1):
                    v = hoja.cell(fila, col + i).value
                    if isinstance(v, (int, float)):
                        acum[clase] = acum.get(clase, 0) + int(v)
        fuera[sentido] = acum
    return fuera


def main() -> None:
    p = argparse.ArgumentParser(description=__doc__,
                                formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("--proyecto", type=int, required=True)
    p.add_argument("--bd", type=Path, default=BD)
    p.add_argument("--manual", type=Path, default=MANUAL)
    p.add_argument("--multiplo", type=float, default=None,
                   help="umbral pesado = multiplo x alto mediano del automovil; "
                        "por omision se calibra en la calzada mejor vista")
    p.add_argument("--holdout", action="store_true",
                   help="calibrar con la primera mitad del horario y validar con "
                        "la segunda; es la unica cifra honesta para la calzada "
                        "donde se calibra, porque ahi el ajuste garantiza el 1.00x")
    a = p.parse_args()

    con = sqlite3.connect(f"file:{a.bd}?mode=ro", uri=True)
    con.row_factory = sqlite3.Row

    cubiertos: set[int] = set()
    for r in con.execute(
        "select v.video_start_time t, v.total_frames f, v.fps from video_jobs v "
        "where v.project_id=? and v.video_start_time is not null "
        "  and (v.status='done' or exists (select 1 from crossings x where x.job_id=v.id))",
        (a.proyecto,),
    ):
        h, m, _ = map(int, r["t"][11:].split(":"))
        cubiertos.update(range(h * 60 + m, h * 60 + m + int((r["f"] or 0) / (r["fps"] or 15) / 60)))
    bins = [b for b in range(0, 24 * 60, 15) if all(x in cubiertos for x in range(b, b + 15))]
    if not bins:
        sys.exit("El video no cubre ningun cuarto de hora completo.")
    print(f"Ventana: {bins[0]//60:02d}:{bins[0]%60:02d} a "
          f"{(bins[-1]+15)//60:02d}:{(bins[-1]+15)%60:02d}  ({len(bins)} cuartos de hora)\n")

    # Nuestros cruces con su alto, por calzada
    if a.holdout:
        # Calibrar con una mitad del horario y medir con la otra. Sin esto,
        # la calzada donde se calibra sale siempre en 1.00x por
        # construccion, y eso no valida nada.
        corte = bins[len(bins) // 2]
        bins_cal = [b for b in bins if b < corte]
        bins_val = [b for b in bins if b >= corte]
        print(f"Holdout: se calibra con {bins_cal[0]//60:02d}:{bins_cal[0]%60:02d}-"
              f"{corte//60:02d}:{corte%60:02d} y se valida con "
              f"{corte//60:02d}:{corte%60:02d}-{(bins_val[-1]+15)//60:02d}:"
              f"{(bins_val[-1]+15)%60:02d}\n")
    else:
        bins_cal = bins_val = bins

    def cruces(rango):
        fuera: dict[str, list[tuple[str, float]]] = {}
        lo, hi = rango[0], rango[-1] + 15
        for r in con.execute(
            "select z.name zona, x.vehicle_type tipo, x.bbox_height alto, x.timestamp ts "
            "from crossings x join video_jobs v on v.id=x.job_id "
            "join zones z on z.id=x.zone_id "
            "where v.project_id=? and x.bbox_height is not null",
            (a.proyecto,),
        ):
            minuto = int(r["ts"][11:13]) * 60 + int(r["ts"][14:16])
            if lo <= minuto < hi:
                fuera.setdefault(r["zona"], []).append((r["tipo"], r["alto"]))
        return fuera

    porzona_cal = cruces(bins_cal)
    porzona = cruces(bins_val)
    con.close()

    manual_cal = leer_manual(a.manual, bins_cal)
    manual = leer_manual(a.manual, bins_val)
    print("Conteo manual en la ventana:")
    for s, cl in manual.items():
        tot = sum(cl.values())
        pes = sum(cl.get(k, 0) for k in PESADAS)
        print(f"  {s}: total {tot}, A {cl.get('A',0)} ({100*cl.get('A',0)/tot:.1f} %), "
              f"pesados {pes} ({100*pes/tot:.1f} %)  " +
              " ".join(f"{k}={cl.get(k,0)}" for k in PESADAS))

    # El emparejamiento calzada-sentido va por VOLUMEN, ambos en orden
    # descendente. Ordenar las calzadas por altura en su lugar las cruzaba
    # al reves —la calzada cercana se emparejaba con el sentido de mas
    # transito, que es el otro— y la validacion salia en 0.71x sin que nada
    # estuviera mal en la clasificacion.
    zonas = sorted(porzona, key=lambda z: -len(porzona[z]))
    sentidos = sorted(manual, key=lambda s: -sum(manual[s].values()))
    # La calibracion del umbral va aparte: se hace en la calzada mejor
    # vista, que es la de alturas mayores, no la de mas transito.
    cercana = max(porzona, key=lambda z: statistics.median(h for _, h in porzona[z]))

    def reparto(zona: str, umbral: float, datos=None) -> tuple[int, int]:
        liv = pes = 0
        for tipo, alto in (datos or porzona)[zona]:
            if tipo in LIVIANAS_YOLO or (tipo == "truck" and alto <= umbral):
                liv += 1
            else:
                pes += 1
        return liv, pes

    mediana_auto = {z: statistics.median(h for t, h in porzona[z] if t == "car")
                    for z in porzona}
    mediana_auto_cal = {z: statistics.median(h for t, h in porzona_cal[z] if t == "car")
                        for z in porzona_cal}

    par = dict(zip(zonas, sentidos))
    if a.multiplo is None:
        # Calibrar: el multiplo que iguala los pesados de la calzada cercana
        objetivo = sum(manual_cal[par[cercana]].get(k, 0) for k in PESADAS)
        mejor, mejor_err = None, None
        for mult in [x / 100 for x in range(100, 401)]:
            _, pes = reparto(cercana, mult * mediana_auto_cal[cercana], porzona_cal)
            err = abs(pes - objetivo)
            if mejor_err is None or err < mejor_err:
                mejor, mejor_err = mult, err
        multiplo = mejor
        print(f"\nUmbral calibrado en {cercana} (la mejor vista): "
              f"{multiplo:.2f} x el alto mediano del automovil")
    else:
        multiplo = a.multiplo
        print(f"\nUmbral impuesto: {multiplo:.2f} x el alto mediano del automovil")

    print()
    print(f'{"calzada":<18}{"umbral":>8}{"":>4}{"livianos":>10}{"manual A":>10}{"razon":>8}'
          f'{"":>3}{"pesados":>9}{"manual":>8}{"razon":>8}')
    for z in zonas:
        s = par[z]
        u = multiplo * mediana_auto[z]
        liv, pes = reparto(z, u)
        mA = manual[s].get("A", 0)
        mP = sum(manual[s].get(k, 0) for k in PESADAS)
        marca = "  <- calibrado aqui" if z == cercana else "  <- validacion independiente"
        print(f"{z:<18}{u:>7.1f}px{'':>4}{liv:>10}{mA:>10}{liv/mA:>7.2f}x"
              f"{'':>3}{pes:>9}{mP:>8}{pes/mP:>7.2f}x{marca}")

    tl = sum(reparto(z, multiplo * mediana_auto[z])[0] for z in zonas)
    tp = sum(reparto(z, multiplo * mediana_auto[z])[1] for z in zonas)
    ml = sum(manual[s].get("A", 0) for s in sentidos)
    mp = sum(sum(manual[s].get(k, 0) for k in PESADAS) for s in sentidos)
    print(f"{'AMBOS SENTIDOS':<18}{'':>8}{'':>4}{tl:>10}{ml:>10}{tl/ml:>7.2f}x"
          f"{'':>3}{tp:>9}{mp:>8}{tp/mp:>7.2f}x")
    print(f"\nComposicion: nuestro {100*tl/(tl+tp):.1f} % livianos / "
          f"{100*tp/(tl+tp):.1f} % pesados")
    print(f"             manual  {100*ml/(ml+mp):.1f} % livianos / "
          f"{100*mp/(ml+mp):.1f} % pesados")

    print("\nSin la correccion, comparando las clases de YOLO en crudo:")
    crudo_liv = sum(1 for z in porzona for t, _ in porzona[z] if t in LIVIANAS_YOLO)
    crudo_pes = sum(1 for z in porzona for t, _ in porzona[z] if t not in LIVIANAS_YOLO)
    print(f"             nuestro {100*crudo_liv/(crudo_liv+crudo_pes):.1f} % livianos / "
          f"{100*crudo_pes/(crudo_liv+crudo_pes):.1f} % pesados  <- la clase `truck` "
          f"de COCO se traga las pickups")


if __name__ == "__main__":
    main()
