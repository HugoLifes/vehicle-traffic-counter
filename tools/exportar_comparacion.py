#!/usr/bin/env python3
"""
Extrae, en un solo JSON, todo lo que el reporte de calibración necesita.

Existe para que el PDF no tenga ni una cifra escrita a mano. La versión
anterior del reporte llevaba la composición vehicular en el código, y eso
contradice lo que el propio documento afirma en su nota metodológica.

Se puede acotar a UNA calzada con `--zona`. Sirve para el caso en que solo
una de las dos se entrega: la cercana a la cámara mide 0.96× y la del fondo
0.84×, así que mezclarlas en una sola cifra la empeora sin necesidad.

Uso, dentro del contenedor del Jetson:
    python3 tools/exportar_comparacion.py --proyecto 2 \
        --zona "Calzada oriente" --salida data/comparacion.json
"""
from __future__ import annotations

import argparse
import json
import re
import sqlite3
import sys
from pathlib import Path

RAIZ = Path(__file__).resolve().parent.parent
MANUAL = RAIZ / "referencias" / "aforo_real" / "conteo_manual_24h.xlsx"

# Clases del conteo manual que son vehículo pesado. "A" incluye pickups y
# camionetas, que es lo que el detector llama `truck` sin serlo.
PESADAS = ("B", "C", "T-S", "T-S-R")

# Qué sentido del aforo manual corresponde a cada calzada. El
# emparejamiento se estableció por correlación del perfil horario
# (r = +1.00), no por el nombre: "Calzada poniente" es ambiguo.
SENTIDO_DE_ZONA = {
    "Calzada oriente": "ote-pte",
    "Calzada poniente": "pte-ote",
}

# Una hora se declara medible cuando la confianza del detector la respalda
# y el conteo no se desploma. El corte de confianza cae en el hueco medido
# entre 0.58 (noche) y 0.67 (día).
CONFIANZA_MINIMA = 0.65
RAZON_MINIMA = 0.75


def leer_manual(ruta: Path, sentidos: set[str]) -> dict:
    """{hora: {'A': n, 'PES': n, 'TOT': n}} sumando los sentidos pedidos."""
    import openpyxl

    libro = openpyxl.load_workbook(str(ruta), data_only=True)
    fuera: dict[int, dict[str, int]] = {}
    for hoja in libro.worksheets:
        cols, sentido = [], None
        for fila in hoja.iter_rows(min_row=1, max_row=6):
            for celda in fila:
                txt = str(celda.value or "").strip()
                if txt.lower().startswith("hr/mov"):
                    cols.append(celda.column)
                elif re.fullmatch(r"(OTE|PTE)-(OTE|PTE)", txt, re.I):
                    sentido = txt.lower()
        if not cols or sentido not in sentidos:
            continue
        # El bloque AM y el PM van lado a lado y repiten las etiquetas de
        # las 12; sin este control se contarían dos veces.
        vistos: set[tuple[int, int]] = set()
        for f in range(1, hoja.max_row + 1):
            for col in cols:
                m = re.match(r"\s*(\d{1,2}):(\d{2})\s*(AM|PM)",
                             str(hoja.cell(f, col).value or ""), re.I)
                if not m:
                    continue
                h, mi, ap = int(m.group(1)), int(m.group(2)), m.group(3).upper()
                if ap == "PM" and h != 12:
                    h += 12
                if ap == "AM" and h == 12:
                    h = 0
                if (h, mi) in vistos:
                    continue
                vistos.add((h, mi))
                d = fuera.setdefault(h, {"A": 0, "PES": 0, "TOT": 0})
                for i, clase in enumerate(("A",) + PESADAS, start=1):
                    v = hoja.cell(f, col + i).value
                    if isinstance(v, (int, float)):
                        d["A" if clase == "A" else "PES"] += int(v)
                        d["TOT"] += int(v)
    return fuera


def main() -> None:
    p = argparse.ArgumentParser(description=__doc__,
                                formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("--proyecto", type=int, required=True)
    p.add_argument("--zona", action="append",
                   help="acotar a esta calzada; repetible. Por omisión, todas")
    p.add_argument("--bd", type=Path, default=RAIZ / "data" / "traffic.db")
    p.add_argument("--manual", type=Path, default=MANUAL)
    p.add_argument("--salida", type=Path, required=True)
    a = p.parse_args()

    con = sqlite3.connect(f"file:{a.bd}?mode=ro", uri=True)
    con.row_factory = sqlite3.Row

    zonas = a.zona or [r["name"] for r in con.execute(
        "select name from zones where project_id=? and kind='calzada'", (a.proyecto,))]
    desconocidas = [z for z in zonas if z not in SENTIDO_DE_ZONA]
    if desconocidas:
        sys.exit(f"No sé qué sentido del aforo manual corresponde a {desconocidas}. "
                 f"Agrégalo a SENTIDO_DE_ZONA.")
    sentidos = {SENTIDO_DE_ZONA[z] for z in zonas}
    manual = leer_manual(a.manual, sentidos)

    # Solo horas con las 6 grabaciones contadas: media hora procesada
    # compararía contra una hora entera de aforo manual.
    completas = {r[0] for r in con.execute(
        """select substr(video_start_time,12,2) h from video_jobs
           where project_id=? and status='done'
           group by h having count(*)=6""", (a.proyecto,))}

    marcas = ",".join("?" * len(zonas))
    horas = []
    for r in con.execute(
        f"""select substr(v.video_start_time,12,2) h, count(*) n,
                   avg(x.confidence) cf
            from crossings x join video_jobs v on v.id=x.job_id
            join zones z on z.id=x.zone_id
            where v.project_id=? and z.name in ({marcas})
            group by h order by h""", (a.proyecto, *zonas)):
        if r["h"] not in completas:
            continue
        real = manual.get(int(r["h"]), {}).get("TOT", 0)
        if not real:
            continue
        razon = r["n"] / real
        horas.append({
            "hora": f"{int(r['h']):02d}:00",
            "nuestro": r["n"],
            "manual": real,
            "razon": round(razon, 3),
            "confianza": round(r["cf"], 3),
            "medible": bool(r["cf"] >= CONFIANZA_MINIMA and razon >= RAZON_MINIMA),
        })
    con.close()

    medibles = [h for h in horas if h["medible"]]
    if not medibles:
        sys.exit("Ninguna hora resultó medible con estos datos.")
    h0 = int(medibles[0]["hora"][:2])
    h1 = int(medibles[-1]["hora"][:2]) + 1

    # La composición se saca por la misma vía que la usa la plataforma, para
    # que el informe no pueda decir una cosa y la pantalla otra.
    sys.path.insert(0, str(RAIZ))
    from src.storage import traffic_db

    reporte = traffic_db.get_interval_counts(a.proyecto, 15)

    # Que carriles pertenecen a las calzadas pedidas. El carril no trae el
    # nombre de su zona, asi que se resuelve por lane_configs.zone_id.
    con2 = sqlite3.connect(f"file:{a.bd}?mode=ro", uri=True)
    con2.row_factory = sqlite3.Row
    marcas_z = ",".join("?" * len(zonas))
    carriles_ok = {r["id"] for r in con2.execute(
        f"""select l.id from lane_configs l join zones z on z.id = l.zone_id
            where l.project_id = ? and l.active = 1 and z.name in ({marcas_z})""",
        (a.proyecto, *zonas))}
    con2.close()
    if not carriles_ok:
        sys.exit(f"Ninguna linea activa esta atada a {zonas}.")

    comp = {"A": 0, "PESADO": 0, "SIN_RESOLVER": 0}
    niveles = []
    for carril in reporte["lanes"]:
        if carril["lane_id"] not in carriles_ok:
            continue
        for iv in carril["intervals"]:
            hh = int(iv["start"][11:13])
            if not (h0 <= hh < h1):
                continue
            for t, cn in iv.get("by_vehicle_type", {}).items():
                comp[t] = comp.get(t, 0) + cn["in"] + cn["out"]
        niveles.append(carril.get("nivel_clasificacion"))

    m_A = sum(manual.get(h, {}).get("A", 0) for h in range(h0, h1))
    m_P = sum(manual.get(h, {}).get("PES", 0) for h in range(h0, h1))
    n_cl = comp["A"] + comp["PESADO"]

    salida = {
        "zonas": zonas,
        "sentidos": sorted(sentidos),
        "horario_medible": {"desde": f"{h0:02d}:00", "hasta": f"{h1:02d}:00",
                            "horas": len(medibles)},
        "totales_medibles": {
            "nuestro": sum(h["nuestro"] for h in medibles),
            "manual": sum(h["manual"] for h in medibles),
            "razon": round(sum(h["nuestro"] for h in medibles)
                           / sum(h["manual"] for h in medibles), 3),
            "razon_min": min(h["razon"] for h in medibles),
            "razon_max": max(h["razon"] for h in medibles),
        },
        "composicion": {
            "nuestro_A_pct": round(100 * comp["A"] / n_cl, 1) if n_cl else None,
            "nuestro_PES_pct": round(100 * comp["PESADO"] / n_cl, 1) if n_cl else None,
            "manual_A_pct": round(100 * m_A / (m_A + m_P), 1) if (m_A + m_P) else None,
            "manual_PES_pct": round(100 * m_P / (m_A + m_P), 1) if (m_A + m_P) else None,
            "clasificados_pct": round(100 * n_cl / (n_cl + comp["SIN_RESOLVER"]), 1)
            if (n_cl + comp["SIN_RESOLVER"]) else None,
            "niveles": sorted(set(n for n in niveles if n)),
        },
        "horas": horas,
    }
    a.salida.parent.mkdir(parents=True, exist_ok=True)
    a.salida.write_text(json.dumps(salida, ensure_ascii=False, indent=2),
                        encoding="utf-8")

    t = salida["totales_medibles"]
    print(f"{a.salida}")
    print(f"  calzadas: {', '.join(zonas)}")
    print(f"  horario medible: {h0:02d}:00 a {h1:02d}:00 ({len(medibles)} horas)")
    print(f"  razon: {t['razon']:.3f}x  (por hora entre {t['razon_min']:.2f} "
          f"y {t['razon_max']:.2f})")
    c = salida["composicion"]
    print(f"  composicion: {c['nuestro_A_pct']} / {c['nuestro_PES_pct']} %  "
          f"contra {c['manual_A_pct']} / {c['manual_PES_pct']} %")


if __name__ == "__main__":
    main()
